"""
Extract the ranking table ("Tabella A") from the GSE PDF
(Graduatoria_DM_FERXtransitorio_PC_EOL_2025_TabA.pdf) and save it as an .xlsx file.

Usage:
    python extract_table.py input.pdf output.xlsx
"""

import sys
import pdfplumber
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


def extract_rows(pdf_path):
    """Extract table rows from every page of the PDF using pdfplumber."""
    all_rows = []
    header = None

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            # pdfplumber's table finder works well on GSE's ruled tables
            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    # Clean up whitespace / None cells
                    clean_row = [
                        (cell.replace("\n", " ").strip() if cell else "")
                        for cell in row
                    ]
                    if not any(clean_row):
                        continue
                    if header is None:
                        header = clean_row
                    elif clean_row == header:
                        continue  # repeated header on later pages
                    else:
                        all_rows.append(clean_row)

    return header, all_rows


def main(pdf_path, xlsx_path):
    header, rows = extract_rows(pdf_path)

    if header is None:
        raise RuntimeError("No table found in the PDF.")

    df = pd.DataFrame(rows, columns=header)

    # Write to Excel
    df.to_excel(xlsx_path, index=False, sheet_name="Tabella A")

    # Light formatting pass
    wb = load_workbook(xlsx_path)
    ws = wb["Tabella A"]

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    body_font = Font(name="Arial", size=10)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center")

    # Auto-size columns (rough heuristic, capped)
    for col_idx, col_cells in enumerate(ws.columns, start=1):
        max_len = max((len(str(c.value)) if c.value else 0) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)

    ws.freeze_panes = "A2"
    wb.save(xlsx_path)
    print(f"Saved {len(df)} rows x {len(df.columns)} columns to {xlsx_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python extract_table.py input.pdf output.xlsx")
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])
